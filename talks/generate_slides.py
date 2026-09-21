"""Generate the GDG Paris talk .pptx from slides_content.py, in GroupBees colors.

Usage (from the repo root):
    uv run --with python-pptx python talks/generate_slides.py

Output: talks/slides/talk_gdg_paris_adk.pptx

Design: GroupBees identity — cyan → purple gradient, near-black ink, white.
Assets in talks/assets/ (logo + mascot downloaded from groupbees.fr, USB-C icon
drawn locally). Slide types: title, bullets, image, code, demo, takeaways, end.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from slides_content import FOOTER_TEXT, SLIDES

ROOT = Path(__file__).parent
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "slides" / "talk_gdg_paris_adk.pptx"

# ── GroupBees palette (sampled from groupbees.fr) ────────────────────────────
INK = RGBColor(0x14, 0x14, 0x1A)        # near-black of the wordmark
CYAN = RGBColor(0x00, 0xB5, 0xFB)       # gradient start
PURPLE = RGBColor(0x97, 0x4E, 0xFF)     # gradient end
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xC9, 0xCF, 0xDC)       # muted text on dark
GRAY = RGBColor(0x5B, 0x63, 0x72)       # muted text on light
BODY_INK = RGBColor(0x2A, 0x2F, 0x3A)   # bullet / body text
CARD = RGBColor(0x1B, 0x1B, 0x24)       # code card
CODE_FG = RGBColor(0xE8, 0xEA, 0xF2)
HEX_TINT = RGBColor(0xEB, 0xF6, 0xFE)   # honeycomb watermark

HEADING_FONT = "Avenir Next"
BODY_FONT = "Avenir Next"
CODE_FONT = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ── low-level helpers ────────────────────────────────────────────────────────
def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _set_run(run, *, size=None, bold=None, italic=None, color=None, font=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = font or BODY_FONT


def _textbox(slide, x, y, w, h, *, anchor=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if anchor is not None:
        tf.vertical_anchor = anchor
    return tf


def _line(tf, text, *, first=False, align=PP_ALIGN.LEFT, space_after=0, **run_kw):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    _set_run(run, **run_kw)
    return p


def _rect(slide, x, y, w, h, color, *, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _gradient(slide, x, y, w, h, *, angle=0.0, shape=MSO_SHAPE.RECTANGLE):
    """Cyan → purple gradient shape — the GroupBees signature accent."""
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.gradient()
    s.fill.gradient_angle = angle
    stops = s.fill.gradient_stops
    stops[0].color.rgb = CYAN
    stops[0].position = 0.0
    stops[1].color.rgb = PURPLE
    stops[1].position = 1.0
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _dark_background(slide):
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, INK)


def _honeycomb(slide, x, y, *, cols=2, rows=3, size=Inches(0.62), color=HEX_TINT):
    """Light honeycomb watermark (a nod to the bees, kept very subtle)."""
    dx = size * 0.86
    dy = size * 0.75
    for r in range(rows):
        for c in range(cols):
            _rect(
                slide,
                int(x + c * dx + (r % 2) * dx / 2),
                int(y + r * dy),
                int(size),
                int(size),
                color,
                shape=MSO_SHAPE.HEXAGON,
            )


def _picture(slide, name, *, left=None, top=None, width=None, height=None):
    path = ASSETS / name
    if not path.exists():
        return None
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def _footer(slide, index, total, *, dark=False):
    """Logo left, talk title + slide number right."""
    if index == 0:
        return
    _picture(slide, "logo_light.png" if dark else "logo.png",
             left=Inches(0.62), top=Inches(6.85), height=Inches(0.3))
    tf = _textbox(slide, Inches(7.5), Inches(6.88), Inches(5.2), Inches(0.35))
    _line(tf, f"{FOOTER_TEXT}   ·   {index}/{total}", first=True,
          align=PP_ALIGN.RIGHT, size=11, color=MIST if dark else GRAY)


def _title(slide, text, *, dark=False):
    tf = _textbox(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(1.0))
    _line(tf, text, first=True, size=34, bold=True,
          color=WHITE if dark else INK, font=HEADING_FONT)
    _gradient(slide, Inches(0.7), Inches(1.32), Inches(1.5), Inches(0.08))


def _notes(slide, notes):
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _bullet_hex(slide, y, i):
    """Small alternating cyan/purple hexagon used as a bullet marker."""
    color = CYAN if i % 2 == 0 else PURPLE
    _rect(slide, Inches(0.78), y, Inches(0.19), Inches(0.19), color,
          shape=MSO_SHAPE.HEXAGON)


# ── renderers ────────────────────────────────────────────────────────────────
def render_title(prs, data):
    slide = _blank(prs)
    _dark_background(slide)
    _gradient(slide, 0, 0, SLIDE_W, Inches(0.12))
    _picture(slide, "logo_light.png", left=Inches(0.7), top=Inches(0.62),
             width=Inches(2.5))
    _picture(slide, "bee.png", left=Inches(9.4), top=Inches(1.5), height=Inches(4.6))

    tf = _textbox(slide, Inches(0.7), Inches(2.5), Inches(8.6), Inches(1.6))
    _line(tf, data["title"], first=True, size=54, bold=True, color=WHITE,
          font=HEADING_FONT)
    if data.get("subtitle"):
        tf2 = _textbox(slide, Inches(0.7), Inches(4.1), Inches(8.2), Inches(1.2))
        _line(tf2, data["subtitle"], first=True, size=22, color=MIST)
    _gradient(slide, Inches(0.7), Inches(5.15), Inches(2.4), Inches(0.1))
    if data.get("footer"):
        tf3 = _textbox(slide, Inches(0.7), Inches(5.6), Inches(8.2), Inches(0.5))
        _line(tf3, data["footer"], first=True, size=16, color=CYAN)
    _notes(slide, data.get("notes"))


def render_end(prs, data):
    slide = _blank(prs)
    _dark_background(slide)
    _gradient(slide, 0, 0, SLIDE_W, Inches(0.12))
    _picture(slide, "bee.png", left=Inches(9.6), top=Inches(1.8), height=Inches(4.2))

    tf = _textbox(slide, Inches(0.7), Inches(2.4), Inches(8.6), Inches(1.4))
    _line(tf, data["title"], first=True, size=54, bold=True, color=WHITE,
          font=HEADING_FONT)
    _gradient(slide, Inches(0.7), Inches(3.9), Inches(2.4), Inches(0.1))
    tf2 = _textbox(slide, Inches(0.7), Inches(4.3), Inches(8.2), Inches(2.0))
    for i, line in enumerate(data.get("lines", [])):
        _line(tf2, line, first=(i == 0), size=20, color=MIST, space_after=10)
    _notes(slide, data.get("notes"))


def render_bullets(prs, data):
    slide = _blank(prs)
    _honeycomb(slide, Inches(11.45), Inches(1.9))
    _title(slide, data["title"])

    image = data.get("image")
    text_w = Inches(7.7) if image else Inches(11.5)
    if image:
        _picture(slide, image, left=Inches(9.6), top=Inches(2.2),
                 height=Inches(3.2) if image != "usbc.png" else None,
                 width=Inches(3.2) if image == "usbc.png" else None)

    # ~ chars per line at 21pt in the available width; used to reserve room for
    # wrapped bullets so the list never bunches up.
    per_line = int(text_w / Inches(0.132))
    y = Inches(1.95)
    for i, bullet in enumerate(data["bullets"]):
        _bullet_hex(slide, y + Inches(0.12), i)
        tf = _textbox(slide, Inches(1.2), y, text_w, Inches(0.6))
        _line(tf, bullet, first=True, size=21, color=BODY_INK)
        lines = max(1, -(-len(bullet) // per_line))
        y += Inches(0.70) + Inches(0.38) * (lines - 1)

    if data.get("kicker"):
        _gradient(slide, Inches(0.7), Inches(6.2), Inches(11.9), Inches(0.05))
        tf = _textbox(slide, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.5))
        _line(tf, data["kicker"], first=True, size=18, bold=True, color=PURPLE)
    _notes(slide, data.get("notes"))


def render_image(prs, data):
    slide = _blank(prs)
    _title(slide, data["title"])
    path = ASSETS.parent / data["image_path"] if data.get("image_path") else None
    if path and not path.is_absolute():
        path = (ROOT / data["image_path"]).resolve()

    body_top, body_h, body_w = Inches(1.65), Inches(4.9), Inches(11.9)
    if path and path.exists():
        pic = slide.shapes.add_picture(str(path), Inches(0.7), body_top, height=body_h)
        if pic.width > body_w:
            ratio = body_w / pic.width
            pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
        pic.left = int((SLIDE_W - pic.width) / 2)
    else:
        box = _rect(slide, Inches(0.7), body_top, body_w, body_h, HEX_TINT)
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _line(tf, f"[ {data.get('image_placeholder', 'image')} ]", first=True,
              align=PP_ALIGN.CENTER, size=18, italic=True, color=GRAY)
    if data.get("caption"):
        tf = _textbox(slide, Inches(0.7), Inches(6.6), Inches(11.9), Inches(0.4))
        _line(tf, data["caption"], first=True, align=PP_ALIGN.CENTER, size=14,
              italic=True, color=GRAY)
    _notes(slide, data.get("notes"))


def render_code(prs, data):
    slide = _blank(prs)
    _title(slide, data["title"])

    lines = data["code"].split("\n")
    has_kicker = bool(data.get("kicker"))
    max_h = 4.35 if has_kicker else 5.05
    card_h = Inches(min(max_h, 0.5 + 0.255 * len(lines)))
    card = _rect(slide, Inches(0.7), Inches(1.62), Inches(11.9), card_h, CARD)
    _gradient(slide, Inches(0.7), Inches(1.62), Inches(0.07), card_h, angle=90.0)

    tf = card.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.35)
    tf.margin_top = Inches(0.22)
    tf.margin_right = tf.margin_bottom = Inches(0.2)
    for i, line in enumerate(lines):
        stripped = line.strip()
        comment = stripped.startswith("#") or stripped.startswith("//")
        _line(tf, line or " ", first=(i == 0), size=14, font=CODE_FONT,
              color=CYAN if comment else CODE_FG)

    if has_kicker:
        tf2 = _textbox(slide, Inches(0.7), Inches(1.62) + card_h + Inches(0.22),
                       Inches(11.9), Inches(0.5))
        _line(tf2, data["kicker"], first=True, size=17, bold=True, color=PURPLE)
    _notes(slide, data.get("notes"))


def render_demo(prs, data):
    slide = _blank(prs)
    _dark_background(slide)
    _gradient(slide, 0, 0, SLIDE_W, Inches(0.12))
    _picture(slide, "bee.png", left=Inches(10.2), top=Inches(2.3), height=Inches(3.6))

    tf = _textbox(slide, Inches(0.8), Inches(1.1), Inches(8.5), Inches(1.5))
    _line(tf, "DÉMO", first=True, size=88, bold=True, color=WHITE, font=HEADING_FONT)
    _gradient(slide, Inches(0.85), Inches(2.55), Inches(3.2), Inches(0.12))

    tf2 = _textbox(slide, Inches(0.85), Inches(2.95), Inches(8.4), Inches(0.8))
    _line(tf2, data["title"], first=True, size=26, color=CYAN)

    y = Inches(3.9)
    for i, step in enumerate(data.get("steps", [])):
        _rect(slide, Inches(0.88), y + Inches(0.1), Inches(0.17), Inches(0.17),
              CYAN if i % 2 == 0 else PURPLE, shape=MSO_SHAPE.HEXAGON)
        tf3 = _textbox(slide, Inches(1.3), y, Inches(8.4), Inches(0.5))
        _line(tf3, step, first=True, size=19, color=MIST)
        y += Inches(0.62)
    _notes(slide, data.get("notes"))


def render_takeaways(prs, data):
    slide = _blank(prs)
    _honeycomb(slide, Inches(11.45), Inches(1.9))
    _title(slide, data["title"])

    y = Inches(2.0)
    for i, (headline, sub) in enumerate(data["items"]):
        badge = _rect(slide, Inches(0.72), y, Inches(0.62), Inches(0.62),
                      CYAN if i % 2 == 0 else PURPLE, shape=MSO_SHAPE.HEXAGON)
        btf = badge.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _line(btf, str(i + 1), first=True, align=PP_ALIGN.CENTER, size=22,
              bold=True, color=WHITE, font=HEADING_FONT)

        tf = _textbox(slide, Inches(1.65), y - Inches(0.02), Inches(10.8), Inches(0.6))
        _line(tf, headline, first=True, size=27, bold=True, color=INK,
              font=HEADING_FONT)
        tf2 = _textbox(slide, Inches(1.65), y + Inches(0.5), Inches(10.8), Inches(0.6))
        _line(tf2, sub, first=True, size=18, color=GRAY)
        y += Inches(1.45)
    _notes(slide, data.get("notes"))


RENDERERS = {
    "title": render_title,
    "bullets": render_bullets,
    "image": render_image,
    "code": render_code,
    "demo": render_demo,
    "takeaways": render_takeaways,
    "end": render_end,
}

DARK_TYPES = {"title", "demo", "end"}


def build_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    total = len(SLIDES)
    for i, data in enumerate(SLIDES):
        renderer = RENDERERS.get(data["type"])
        if renderer is None:
            raise ValueError(f"Slide {i + 1}: unknown type {data['type']!r}")
        renderer(prs, data)
        _footer(prs.slides[i], i, total - 1, dark=data["type"] in DARK_TYPES)
    return prs


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs = build_presentation()
    prs.save(OUTPUT)
    print(f"OK — {len(prs.slides)} slides → {OUTPUT}")


if __name__ == "__main__":
    main()
